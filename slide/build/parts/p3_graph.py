"""Chặng 3 — Knowledge Graph: bảng → đồ thị, triple, mô hình dữ liệu, schema thật, dựng đồ thị, Cypher, cộng đồng, thực hành 3.

Số lấy từ `graph_report.json`, `ingest_report.json`; Cypher và kết quả từ `examples.json:graph`, chọn theo câu benchmark.
Schema theo `src/vre/index/graph_export.py` (NODES, RELS) và data-contract §6.
"""
from __future__ import annotations

from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from kit import C, MX, SIZE, W, Deck, Slide, rgb

SEC = "Chặng 3 · Knowledge Graph"
GR = "graph_report.json"
NODE_LABELS = ["Listing", "PropertyType", "Province", "District", "Ward", "Street", "Project", "Month", "Amenity",
               "Cluster", "MarketStat"]
REL_TYPES = ["HAS_AMENITY", "IN_CLUSTER", "IN_DISTRICT", "IN_PROJECT", "IN_PROVINCE", "IN_WARD", "OF_TYPE",
             "ON_STREET", "POSTED_IN", "STAT_OF"]
CYPHER_HOT = ["IN_WARD|IN_DISTRICT", "MarketStat", "STAT_OF", "IN_DISTRICT", "IN_PROVINCE", "IN_WARD"]


def build(d: Deck) -> None:
    why(d)
    triple(d)
    models(d)
    schema_rules(d)
    schema(d)
    build_graph(d)
    cypher(d)
    communities(d)
    practice(d)


# ---------------------------------------------------------------- helpers
def node(s: Slide, x, y, w, h, label, fill="amber_bg", size=SIZE["small"], bold=True, color="ink"):
    """Hộp node bo góc, chữ căn giữa."""
    sh = s.rect(x, y, w, h, fill=fill, line=None, radius=0.25)
    s.text(0, 0, 0, 0, label, size=size, bold=bold, color=color, align="center", anchor="middle", shape=sh)
    return sh


def arrow(s: Slide, x1, y1, x2, y2, color="muted", dashed=False, width=1.75):
    """Đường thẳng có mũi tên ở đầu cuối (cạnh có hướng)."""
    cn = s.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = rgb(color)
    cn.line.width = Pt(width)
    ln = cn.line._get_or_add_ln()
    if dashed:
        dash = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
        ln.append(dash)
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return cn


def edge_label(s: Slide, cx, cy, text, w=1.45, color="muted"):
    """Nhãn cạnh nền trắng đặt tại (cx, cy) (tâm)."""
    sh = s.text(cx - w / 2, cy - 0.17, w, 0.34, text, size=SIZE["caption"], bold=True, color=color, align="center",
                anchor="middle")
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb("white")
    return sh


def rel_key(t: str) -> str:
    return f"{GR}:relationships.{t}"


def sum_expr(section: str, names: list[str]) -> str:
    return "=" + "+".join(f"{{{GR}:{section}.{n}}}" for n in names)


# ---------------------------------------------------------------- slides
def why(d: Deck) -> None:
    s = d.slide("kg-why", SEC, "Từ bảng sang đồ thị: quan hệ thành cạnh có tên, thay cho phép JOIN", notes=(
        "Chặng ba bắt đầu bằng câu hỏi: dữ liệu đã nằm gọn trong bảng, vậy cần đồ thị để làm gì? Bên trái là dạng "
        "bảng. Tên phường và tên quận lặp lại ở từng dòng, còn quan hệ phường thuộc quận chỉ nằm ngầm trong chuỗi ký "
        "tự. Muốn gom tin theo quận thì phải so khớp chuỗi hoặc JOIN thêm bảng địa danh. Bên phải là cùng thông tin "
        "dạng đồ thị: mỗi phường là một node, nối lên quận bằng cạnh IN_DISTRICT, nhiều tin cùng trỏ vào một node "
        "phường. Đi từ tin lên phường, quận, tỉnh là đi theo cạnh. Gom mọi tin của một quận là đi ngược các cạnh dẫn "
        "vào quận đó. Slide sau định nghĩa chính xác node, cạnh và bộ ba."))
    # Khối 1: bảng phẳng
    s.text(MX, 1.9, 5.6, 0.45, "Bảng phẳng", size=SIZE["body"], bold=True, color="muted")
    rows = [["Tin", "Phường", "Quận"],
            ["Tin A", "Thảo Điền", "Thủ Đức"],
            ["Tin B", "Thảo Điền", "Thủ Đức"],
            ["Tin C", "Dịch Vọng", "Cầu Giấy"]]
    s.table(MX, 2.45, 5.6, rows, col_w=[1.4, 2.2, 2.0], size=SIZE["body"], row_h=0.6)
    s.text(MX, 4.95, 5.6, 0.9, "Tên quận lặp ở từng dòng; quan hệ chỉ nằm trong chuỗi ký tự.",
           size=SIZE["body"], color="text")
    # Khối 2: đồ thị
    gx = 6.6
    s.text(gx, 1.9, 5.7, 0.45, "Đồ thị", size=SIZE["body"], bold=True, color="amber")
    nw, nh = 1.45, 0.6
    tins = [("Tin A", 2.5), ("Tin B", 3.5), ("Tin C", 4.6)]
    for lab, y in tins:
        node(s, gx, y, nw, nh, lab, fill="blue_bg", size=SIZE["body"])
    px, dx = gx + 2.1, gx + 4.25
    node(s, px, 3.0, nw, nh, "Thảo Điền", fill="teal_bg", size=SIZE["body"])
    node(s, px, 4.6, nw, nh, "Dịch Vọng", fill="teal_bg", size=SIZE["body"])
    node(s, dx, 3.0, 1.75, nh, "Thủ Đức", fill="amber_bg", size=SIZE["body"])
    node(s, dx, 4.6, 1.75, nh, "Cầu Giấy", fill="amber_bg", size=SIZE["body"])
    arrow(s, gx + nw, 2.8, px, 3.3)
    arrow(s, gx + nw, 3.8, px, 3.3)
    arrow(s, gx + nw, 4.9, px, 4.9)
    arrow(s, px + nw, 3.3, dx, 3.3)
    arrow(s, px + nw, 4.9, dx, 4.9)
    edge_label(s, (px + nw + dx) / 2, 2.72, "IN_DISTRICT", w=1.35)
    s.text(gx, 2.0 + 3.35, 5.7, 0.5, "Mỗi địa danh là một node duy nhất.", size=SIZE["body"])
    # Khối 3: kết luận
    s.rect(MX, 5.95, W - 2 * MX, 0.75, fill="amber_bg", line=None, radius=0.1)
    s.text(MX + 0.3, 5.95, W - 2 * MX - 0.6, 0.75,
           [[("Truy vấn = đi theo cạnh.  ", {"bold": True, "color": "ink"}),
             ("Gom tin theo quận là đi ngược các cạnh dẫn vào node quận.", {})]],
           size=SIZE["body"], anchor="middle")


def triple(d: Deck) -> None:
    s = d.slide("kg-triple", SEC, "Knowledge Graph ghép tri thức từ các bộ ba (chủ thể)-[quan hệ]->(đối tượng)",
                notes=(
        "Knowledge graph là mạng các thực thể và quan hệ có hướng giữa chúng. Thực thể là node: tin đăng, phường, "
        "quận, loại hình. Quan hệ là cạnh có tên và có hướng. Đơn vị nhỏ nhất là bộ ba gồm chủ thể, quan hệ, đối "
        "tượng, ví dụ tin đăng nằm trong phường. Hình bên phải là một tin minh hoạ, không phải tin thật trong dữ "
        "liệu: một căn hộ ở phường Thảo Điền có hồ bơi. Một dòng bảng tách thành nhiều bộ ba, và các bộ ba nối với "
        "nhau qua node chung, nên từ tin có thể đi tiếp lên quận rồi lên tỉnh. Slide sau bàn chọn mô hình dữ liệu "
        "nào để lưu các bộ ba này."))
    lx, lw = MX, 5.8
    s.text(lx, 1.95, lw, 1.5, [
        [("Thực thể (node)  ", {"bold": True, "color": "ink"}), ("tin đăng, phường, quận, loại hình", {})],
        [("Quan hệ (cạnh)  ", {"bold": True, "color": "ink"}), ("liên kết có tên, có hướng", {})]],
        size=SIZE["body"], spacing=1.3)
    s.rect(lx, 3.55, lw, 1.9, fill="soft", line=None, radius=0.08)
    s.text(lx + 0.3, 3.7, lw - 0.6, 0.4, "Bộ ba (triple)", size=SIZE["small"], bold=True, color="amber")
    parts = [("Chủ thể", "(Tin)", "blue"), ("Quan hệ", "-[:IN_WARD]->", "amber"), ("Đối tượng", "(Thảo Điền)", "teal")]
    cw = (lw - 0.6) / 3
    for i, (lab, val, col) in enumerate(parts):
        x = lx + 0.3 + i * cw
        s.text(x, 4.15, cw, 0.4, lab, size=SIZE["small"], color="muted", align="center")
        s.text(x, 4.6, cw, 0.6, val, size=18, bold=True, color=col, align="center")
    s.text(lx, 5.7, lw, 0.8, "Một dòng bảng tách thành nhiều bộ ba nối nhau qua node chung.", size=SIZE["body"])
    # Minh hoạ: một tin thành node/cạnh
    cx, cy, nw, nh = 8.55, 3.75, 1.9, 0.6
    node(s, cx, cy, nw, nh, "Tin (minh hoạ)", fill="blue_bg")
    around = [("Căn hộ chung cư", 6.6, 1.95, "OF_TYPE", "violet_bg"),
              ("Hồ bơi", 10.6, 1.95, "HAS_AMENITY", "coral_bg"),
              ("Thảo Điền", 6.6, 5.55, "IN_WARD", "teal_bg")]
    for lab, x, y, rel, fill in around:
        node(s, x, y, nw, nh, lab, fill=fill)
        x1 = cx + nw / 2
        y1 = cy if y < cy else cy + nh
        x2 = x + nw / 2
        y2 = y + nh if y < cy else y
        arrow(s, x1, y1, x2, y2)
        edge_label(s, (x1 + x2) / 2, (y1 + y2) / 2, rel, w=1.45)
    node(s, 10.6, 5.55, nw, nh, "Thủ Đức", fill="amber_bg")
    arrow(s, 6.6 + nw, 5.85, 10.6, 5.85)
    edge_label(s, (6.6 + nw + 10.6) / 2, 5.85, "IN_DISTRICT", w=1.45)
    s.caption("Tin và địa danh minh hoạ, không lấy từ dữ liệu; nhãn cạnh là kiểu cạnh thật trong schema.", y=6.55)


def models(d: Deck) -> None:
    s = d.slide("kg-models", SEC, "Chọn property graph: thuộc tính nằm ngay trên node, truy vấn bằng Cypher", notes=(
        "Có ba cách lưu cùng một tri thức. Property graph, như Neo4j, cho node và cạnh mang thuộc tính, nên giá và "
        "diện tích nằm ngay trên node tin đăng, truy vấn bằng Cypher dễ đọc. RDF là chuẩn bộ ba của W3C, mạnh khi "
        "trao đổi và liên kết dữ liệu mở, nhưng mỗi thuộc tính lại là một bộ ba nên mô hình phình ra. SQL là bảng, "
        "dòng, cột. Nói trung thực: với phép tổng hợp trên bảng phẳng thì SQL hay pandas là đủ và nhanh. Chính đáp án "
        "chuẩn của benchmark được tính bằng SQL trên DuckDB. Đồ thị có lợi khi câu hỏi đi nhiều chặng qua quan hệ. "
        "Slide sau là nguyên tắc chọn node, thuộc tính và cạnh."))
    rows = [["Tiêu chí", "Property graph (nhóm chọn)", "RDF", "SQL"],
            ["Mô hình", "Node, cạnh có thuộc tính", "Bộ ba chuẩn W3C", "Bảng, dòng, cột"],
            ["Truy vấn", "Cypher", "SPARQL", "SQL"],
            ["Mạnh ở", "Đi nhiều chặng, ít JOIN", "Trao đổi dữ liệu mở", "Tổng hợp bảng phẳng"],
            ["Ví dụ", "(Tin)-[:IN_WARD]->(Phường)", "<Tin> <inWard> <Phường>", "JOIN listing, ward"]]
    s.table(MX, 1.95, W - 2 * MX, rows, col_w=[1.9, 3.9, 3.3, 3.033], size=18, row_h=0.62, highlight_row=None)
    s.rect(MX, 5.35, W - 2 * MX, 1.05, fill="soft", line=None, radius=0.1)
    s.text(MX + 0.3, 5.35, W - 2 * MX - 0.6, 1.05,
           [[("Khi nào SQL/pandas đủ:  ", {"bold": True, "color": "ink"}),
             ("phép tổng hợp trên một bảng phẳng. Đáp án chuẩn của benchmark được tính bằng SQL trên DuckDB.", {})]],
           size=SIZE["small"], anchor="middle")


def schema_rules(d: Deck) -> None:
    s = d.slide("kg-schema-rules", SEC, "Thứ cần duyệt qua hoặc gom nhóm là node; thứ để lọc và tính là thuộc tính",
                notes=(
        "Schema là bản thiết kế của đồ thị: nhãn node nào, kiểu cạnh nào, thuộc tính nào được phép có. Câu hỏi quyết "
        "định chỉ có một: có cần duyệt qua hoặc gom nhóm theo đối tượng này không? Nếu có thì nó là node, như phường, "
        "quận, đường, dự án, vì câu hỏi gom tin theo các đơn vị đó. Nếu chỉ để lọc, sắp xếp hay tính toán thì nó là "
        "thuộc tính, như giá, diện tích, số phòng. Cạnh là quan hệ có hướng, tên là động từ hoặc giới từ. Bản này còn "
        "thêm node thống kê tính sẵn MarketStat, để câu tổng hợp đọc một node thay vì quét hàng triệu tin. Slide sau "
        "là schema thật."))
    cards = [("Node", "Đối tượng có danh tính, cần duyệt qua hoặc gom nhóm.", "Ward, District, Street, Project", "teal"),
             ("Thuộc tính", "Giá trị mô tả, dùng để lọc, sắp xếp, tính toán.", "price_vnd, area_m2, bedrooms", "blue"),
             ("Cạnh", "Quan hệ có hướng giữa hai node, tên là động từ hoặc giới từ.", "(Listing)-[:IN_WARD]->(Ward)",
              "amber")]
    cw = (W - 2 * MX - 2 * 0.35) / 3
    for i, (t, body, ex, col) in enumerate(cards):
        x = MX + i * (cw + 0.35)
        s.rect(x, 1.95, cw, 3.4, fill=f"{col}_bg", line=None, radius=0.06)
        s.text(x + 0.3, 2.15, cw - 0.6, 0.55, t, size=26, bold=True, color=col)
        s.text(x + 0.3, 2.85, cw - 0.6, 1.5, body, size=SIZE["body"], color="ink")
        s.text(x + 0.3, 4.4, cw - 0.6, 0.8, ex, size=SIZE["body"], color="text", font="Consolas")
    s.rect(MX, 5.65, W - 2 * MX, 0.85, fill="soft", line=None, radius=0.1)
    s.text(MX + 0.3, 5.65, W - 2 * MX - 0.6, 0.85,
           [[("Tiêu chí:  ", {"bold": True, "color": "ink"}),
             ("có cần đi qua hay gom nhóm theo nó không? Có thì là node, không thì là thuộc tính.", {})]],
           size=SIZE["body"], anchor="middle")


def schema(d: Deck) -> None:
    miss = d.N("ingest_report.json:missing_pct.ward", "dec1")
    s = d.slide("kg-schema", SEC, "Schema thật: Listing ở tâm, cây địa danh, cụm phường và thống kê tính sẵn", notes=(
        f"Đây là schema đang chạy trên toàn bộ dữ liệu: {d.N(GR + ':nodes|len')} nhãn "
        f"node và {d.N(GR + ':relationships|len')} kiểu cạnh. Listing ở giữa, nối sang loại hình, tháng đăng, tiện "
        "ích; bên phải là phường, đường, dự án, đều đi lên quận rồi lên tỉnh. Khoá của phường gồm cả tỉnh và quận, vì "
        "cùng một tên phường có ở nhiều quận; chỉ dùng tên thì các phường trùng tên bị gộp sai. Phường nối vào cụm "
        "phường do Leiden tìm ra. MarketStat là thống kê giá tính sẵn theo khu vực, loại hình và tháng, trỏ vào khu "
        f"vực bằng STAT_OF. Dữ liệu thiếu phường ở {miss}% tin, nên tin đó nối thẳng vào quận, thiếu cả quận thì nối "
        "vào tỉnh. Slide sau: đồ thị được dựng thế nào."))
    nw, nh = 1.85, 0.55
    xa, xb, xc, xd = MX, 3.95, 7.45, 10.85
    r1, r2, r3, r4 = 1.95, 3.15, 4.35, 5.4
    mid = nh / 2
    node(s, xb, r2, nw, nh, "Listing", fill="blue", color="white")
    left = [("PropertyType", r1, "OF_TYPE"), ("Month", r2, "POSTED_IN"), ("Amenity", r3, "HAS_AMENITY")]
    for lab, y, rel in left:
        node(s, xa, y, nw, nh, lab, fill="violet_bg")
        arrow(s, xb, r2 + mid, xa + nw, y + mid)
        edge_label(s, (xb + xa + nw) / 2, (r2 + y) / 2 + mid, rel, w=1.45)
    right = [("Ward", r1, "IN_WARD"), ("Street", r2, "ON_STREET"), ("Project", r3, "IN_PROJECT")]
    for lab, y, rel in right:
        node(s, xc, y, nw, nh, lab, fill="teal_bg")
        arrow(s, xb + nw, r2 + mid, xc, y + mid)
        edge_label(s, (xb + nw + xc) / 2, (r2 + y) / 2 + mid, rel, w=1.45)
    node(s, xd, r1, nw, nh, "Cluster", fill="coral_bg")
    node(s, xd, r2, nw, nh, "District", fill="amber_bg")
    node(s, xd, r3, nw, nh, "Province", fill="amber_bg")
    gap_c = (xc + nw + xd) / 2
    arrow(s, xc + nw, r1 + mid, xd, r1 + mid)
    edge_label(s, gap_c, r1 + mid, "IN_CLUSTER", w=1.4)
    arrow(s, xc + nw, r1 + mid, xd, r2 + mid)
    arrow(s, xc + nw, r2 + mid, xd, r2 + mid)
    edge_label(s, gap_c, r2 + mid, "IN_DISTRICT", w=1.4)
    arrow(s, xc + nw, r3 + mid, xd, r2 + mid)
    arrow(s, xd + nw / 2, r2 + nh, xd + nw / 2, r3)
    edge_label(s, xd + nw / 2, (r2 + nh + r3) / 2, "IN_PROVINCE", w=1.4)
    node(s, xc, r4, nw, nh, "MarketStat", fill="soft", color="ink")
    arrow(s, xc + nw / 2, r4, xc + nw / 2, r3 + nh)
    edge_label(s, xc + nw / 2, (r4 + r3 + nh) / 2, "STAT_OF", w=1.15)
    s.text(9.5, r4, 3.2, 0.75, "STAT_OF còn trỏ tới Ward, District, Province", size=SIZE["caption"],
           color="muted")
    s.rect(MX, r4 - 0.05, 6.45, 1.2, fill="amber_bg", line=None, radius=0.08)
    s.text(MX + 0.2, r4, 6.1, 1.1, [
        [("Cạnh dự phòng  ", {"bold": True, "color": "ink"}), (f"{miss}% tin thiếu phường:", {})],
        [("Listing -[:IN_DISTRICT]-> District; thiếu cả quận: -[:IN_PROVINCE]-> Province", {"font": "Consolas"})]],
        size=SIZE["small"], anchor="middle")
    s.caption("Mọi mũi tên vào District là IN_DISTRICT (từ Ward, Street, Project). Nguồn: graph_export.py.", y=6.6)


def build_graph(d: Deck) -> None:
    s = d.slide("kg-build", SEC, "Dựng đồ thị bằng luật từ cột có cấu trúc, nạp hàng loạt bằng neo4j-admin", notes=(
        "Mọi node và cạnh đều sinh từ cột có cấu trúc bằng luật, không gọi LLM. DuckDB đọc bảng sạch, tính sẵn các "
        "tổng hợp rồi xuất CSV theo định dạng của neo4j-admin. Lệnh database import full nạp toàn bộ đồ thị trong "
        f"{d.N(GR + ':import_seconds', 'dec1')} giây, nhanh hơn nhiều so với chèn từng lô bằng Cypher. Sau đó tạo "
        "ràng buộc khoá, index, rồi kiểm hai bất biến: số tin và tổng giá phải khớp bảng sạch. Phần khó là entity "
        "resolution: cùng một quận có thể ghi là chữ số, viết tắt hay đầy đủ, nếu không chuẩn hoá sẽ thành nhiều node. "
        "Tiện ích lấy bằng từ điển biểu thức chính quy trên mô tả; MarketStat là trung vị giá tính bằng DuckDB, chỉ "
        "tạo khi nhóm đủ số tin tối thiểu. Slide sau: truy vấn đồ thị bằng Cypher."))
    s.flow([("DuckDB", "bảng sạch, tổng hợp sẵn, xuất CSV"),
            (d.A("neo4j-admin", "tên công cụ"), "database import full"),
            ("Ràng buộc, index", "khoá duy nhất mỗi nhãn"),
            ("Kiểm bất biến", "số tin, tổng giá khớp bảng sạch")], y=1.95, h=1.35,
           fills=["blue_bg", "amber_bg", "teal_bg", "violet_bg"])
    s.stat(MX, 3.65, 3.3, f"{d.N(GR + ':import_seconds', 'dec1')} giây", "thời gian import full toàn bộ đồ thị",
           accent="amber", h=1.45)
    s.stat(MX, 5.25, 3.3, d.N(GR + ':nodes.Listing'), "node Listing, mỗi tin một node",
           accent="blue", h=1.4)
    # Entity resolution
    ex, ew = 4.3, 4.25
    s.text(ex, 3.65, ew, 0.45, "Entity resolution địa danh", size=SIZE["small"], bold=True, color="ink")
    pairs = [(d.A("“Q.7”, “q7”, “7”", "ví dụ chuẩn hoá trong places.normalize_place"),
              d.A("Quận 7", "ví dụ chuẩn hoá trong places.normalize_place")),
             (d.A("“P.5”", "ví dụ chuẩn hoá trong places.normalize_place"),
              d.A("Phường 5", "ví dụ chuẩn hoá trong places.normalize_place")),
             ("“TP. Hồ Chí Minh”", "Hồ Chí Minh")]
    for i, (a, b) in enumerate(pairs):
        y = 4.2 + i * 0.75
        s.rect(ex, y, 1.95, 0.6, fill="soft", line=None, radius=0.15)
        s.text(ex, y, 1.95, 0.6, a, size=SIZE["small"], align="center", anchor="middle")
        s.rect(ex + 2.0, y + 0.2, 0.3, 0.2, fill="muted", line=None, shape=_right_arrow())
        s.rect(ex + 2.35, y, 1.9, 0.6, fill="teal_bg", line=None, radius=0.15)
        s.text(ex + 2.35, y, 1.9, 0.6, b, size=SIZE["small"], bold=True, align="center", anchor="middle")
    # Tổng hợp + tiện ích
    tx, tw = 8.95, W - MX - 8.95
    s.card(tx, 3.65, tw, 1.45, "MarketStat", "trung vị giá/m² theo khu vực, loại hình, tháng (DuckDB)",
           accent="amber")
    s.card(tx, 5.25, tw, 1.4, "Amenity", "từ điển regex trên mô tả, không dùng LLM", accent="coral")


def _right_arrow():
    from pptx.enum.shapes import MSO_SHAPE
    return MSO_SHAPE.RIGHT_ARROW


def cypher(d: Deck) -> None:
    """Một khối Cypher mỗi slide; ví dụ chọn theo câu benchmark trong examples.json (không theo thứ tự)."""
    cases = [
        ("kg-cypher-agg", "AG01", "Tổng hợp", "amber",
         "Câu tổng hợp đọc node MarketStat tính sẵn thay vì quét hàng triệu tin",
         "Câu hỏi tổng hợp như giá trung vị theo quận không cần quét hàng triệu tin lúc hỏi. Khi dựng đồ thị, DuckDB đã "
         "tính sẵn trung vị theo khu vực, loại hình và tháng, lưu thành node MarketStat. Mô hình chỉ cần viết một MATCH "
         "đến đúng node đó. Phần tô màu là điều kiện chọn khu vực và loại hình. Dưới khối code là số dòng trả về và độ "
         "trễ đo trong lượt chạy thật. Slide sau: câu đa chặng."),
        ("kg-cypher-multihop", "MH01", "Đa chặng", "teal",
         "Câu đa chặng đi qua nhiều cạnh liên tiếp, việc SQL cần nhiều phép JOIN",
         "Câu đa chặng phải đi từ dự án sang quận, rồi từ quận sang tỉnh, sau đó so sánh với thống kê của cả tỉnh. "
         "Trong Cypher, mỗi bước là một cạnh trong cùng một mẫu MATCH; SQL tương đương cần nhiều phép JOIN. Phần tô "
         "màu là các cạnh được đi qua. Đây là câu Cypher mô hình sinh thật trong lượt eval, chạy trên Neo4j với giao "
         "dịch chỉ đọc. Slide sau chuyển sang thuật toán trên đồ thị."),
    ]
    for name, qid, lab, col, title, notes in cases:
        k = d.pick("examples.json:graph", qid)
        s = d.slide(name, SEC, title, notes=notes)
        s.rect(MX, 1.95, 0.08, 0.42, fill=col, line=None)
        s.text(MX + 0.18, 1.93, 3.0, 0.45, lab, size=SIZE["body"], bold=True, color=col)
        q = d.T(f"{k}.question") if k else d.A(d.Q(qid)["question"], f"đề bài {qid}")
        s.text(MX + 2.0, 1.97, W - 2 * MX - 2.0, 0.6, "“" + q + "”", size=SIZE["small"], color="ink")
        s.code(MX, 2.75, W - 2 * MX, 3.25, d.T(f"{k}.cypher") if k else d.A("⟨?⟩", "chưa có ví dụ"),
               highlight=CYPHER_HOT, size=SIZE["small"])
        if k:
            s.text(MX, 6.15, W - 2 * MX, 0.45, [[(f"{d.N(k + '.rows|len')} dòng kết quả", {"bold": True, "color": "ink"}),
                   (f"  ·  {d.N(k + '.latency_ms', 'int')} ms trên Neo4j", {})]], size=SIZE["small"])


def communities(d: Deck) -> None:
    c = f"{GR}:communities"
    s = d.slide("kg-communities", SEC, "Leiden gom phường có hồ sơ giá giống nhau thành cụm, đo bằng modularity",
                notes=(
        "Hai khái niệm. Centrality đo mức quan trọng của một node theo vị trí, dạng đơn giản nhất là bậc chia cho số "
        "node còn lại; ở đây chỉ để hiểu, pipeline không dùng nó. Modularity đo các cụm dày cạnh hơn ngẫu nhiên bao "
        "nhiêu. Điểm quan trọng: không chạy Leiden trên cây tỉnh, quận, phường, vì cây không có cạnh chéo nên không "
        "có cụm. Code dựng đồ thị tương đồng: mỗi phường đủ số tin là một hồ sơ gồm log giá trên mét vuông trung vị, "
        "log số tin và tỷ lệ từng loại hình, chuẩn hoá rồi nối với "
        f"{d.N(c + '.k')} phường gần nhất. Leiden với seed {d.N(c + '.seed')} chia {d.N(c + '.n_wards')} phường "
        f"thành {d.N(c + '.n_clusters')} cụm, modularity {d.N(c + '.modularity', 'dec2')}, tức cấu trúc cụm rõ. "
        "Mỗi phường được nối vào node Cluster. Slide sau là con số kiểm tra toàn đồ thị."))
    lw = 6.2
    s.text(MX, 1.95, lw, 0.4, "Độ trung tâm bậc (degree centrality)", size=SIZE["small"],
           bold=True, color="muted")
    s.text(MX, 2.35, lw, 0.6, d.A("C(v) = deg(v) / (n − 1)", "công thức lý thuyết"), size=24, bold=True,
           color="ink", font="Cambria Math")
    s.text(MX, 3.15, lw, 0.4, "Modularity", size=SIZE["small"],
           bold=True, color="muted")
    s.text(MX, 3.55, lw, 0.6, d.A("Q = Σc [ Lc / m − (dc / 2m)² ]", "công thức lý thuyết"), size=24, bold=True,
           color="ink", font="Cambria Math")
    s.text(MX, 4.3, lw, 1.0, d.A("Ví dụ: hai tam giác nối một cạnh, m = 7, Q = 5/14 ≈ 0,357; node cầu C = 0,6",
                                  "ví dụ minh hoạ công thức"), size=SIZE["small"], color="text")
    s.rect(MX, 5.3, lw, 1.15, fill="soft", line=None, radius=0.08)
    s.text(MX + 0.25, 5.3, lw - 0.5, 1.15,
           [[("Đồ thị Leiden chạy trên  ", {"bold": True, "color": "ink"}),
             ("kNN giữa các phường theo hồ sơ giá/m², số tin, cơ cấu loại hình; không phải cây hành chính.", {})]],
           size=SIZE["small"], anchor="middle")
    rx = 7.25
    sw = (W - MX - rx - 0.3) / 2
    s.stat(rx, 1.95, sw, d.N(c + ".n_wards"), "phường đủ tin vào đồ thị kNN", accent="teal")
    s.stat(rx + sw + 0.3, 1.95, sw, d.N(c + ".n_clusters"), "cụm phường (node Cluster)", accent="amber")
    s.stat(rx, 3.7, sw, d.N(c + ".modularity", "dec2"), "modularity Q", accent="violet")
    s.stat(rx + sw + 0.3, 3.7, sw, f"k = {d.N(c + '.k')}", f"láng giềng gần nhất, seed {d.N(c + '.seed')}",
           accent="blue")
    s.text(rx, 5.45, W - MX - rx, 1.0, "Chỉ số centrality không được tính trong pipeline; chỉ modularity.",
           size=SIZE["small"], color="muted")


def practice(d: Deck) -> None:
    nodes = NODE_LABELS
    inv = d.V(f"{GR}:invariants") or {}
    all_ok = bool(inv) and all(v is True for v in inv.values())
    s = d.slide("kg-practice-3", SEC, "Thực hành 3: đồ thị Neo4j trên toàn bộ dữ liệu, bất biến khớp bảng sạch", notes=(
        f"Đồ thị nạp từ đúng bảng sạch của chặng một: tổng cộng {d.N(sum_expr('nodes', NODE_LABELS))} node và "
        f"{d.N(sum_expr('relationships', REL_TYPES))} cạnh. Listing chiếm phần lớn node với "
        f"{d.N(GR + ':nodes.Listing')} tin; MarketStat có {d.N(GR + ':nodes.MarketStat')} node thống kê. Bên cạnh, "
        f"IN_WARD có {d.N(rel_key('IN_WARD'))} cạnh, ít hơn số tin vì tin thiếu phường đi qua IN_DISTRICT. "
        "HAS_AMENITY nhiều hơn số tin vì một tin có thể có nhiều tiện ích. Hai bất biến được kiểm tự động sau import: "
        "số node Listing bằng số tin sạch, và tổng giá khớp. " + ("Cả hai đều đạt. " if all_ok else "Có bất biến chưa đạt, cần xem lại. ") + "Hết phần knowledge graph; tiếp theo là "
        "ghép RAG và đồ thị thành GraphRAG và đo xem có đáng không."))
    half = (W - 2 * MX - 0.4) / 2
    s.text(MX, 1.85, half, 0.4, "Số node theo nhãn", size=SIZE["small"], bold=True, color="ink")
    s.chart(MX, 2.3, half, 3.65, "bar", list(reversed(nodes)),
            [("Số node", [f"{GR}:nodes.{n}" for n in reversed(nodes)], "amber")], number_format="#,##0")
    x2 = MX + half + 0.4
    s.text(x2, 1.85, half, 0.4, "Số cạnh theo kiểu", size=SIZE["small"], bold=True, color="ink")
    s.chart(x2, 2.3, half, 3.65, "bar", list(reversed(REL_TYPES)),
            [("Số cạnh", [rel_key(t) for t in reversed(REL_TYPES)], "teal")], number_format="#,##0")
    ok = lambda k: ("đạt", "teal") if inv.get(k) is True else ("KHÔNG đạt", "coral")  # noqa: E731
    (t1, c1), (t2, c2) = ok("listing_count_ok"), ok("price_sum_ok")
    s.rect(MX, 6.05, W - 2 * MX, 0.6, fill="soft", line=None, radius=0.15)
    s.text(MX + 0.25, 6.05, W - 2 * MX - 0.5, 0.6, [[
        ("Tổng  ", {"bold": True, "color": "ink"}),
        (f"{d.N(sum_expr('nodes', NODE_LABELS))} node · {d.N(sum_expr('relationships', REL_TYPES))} cạnh", {}),
        ("     Số Listing = số tin sạch: ", {}), (t1, {"bold": True, "color": c1}),
        ("     Tổng giá khớp: ", {}), (t2, {"bold": True, "color": c2})]], size=SIZE["small"], anchor="middle")
