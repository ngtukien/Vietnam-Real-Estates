"""Mọi con số trên slide phải truy được về results/*.json.

  extract <pptx>                    liệt kê mọi số (thân, bảng, notes, chart)
  check   <pptx> --map M [--slides b06,n-eda-1]
  apply   <pptx> --map M [--out X]  ghi giá trị từ JSON vào slot của từng mục map, cập nhật data chart
  html    <html> --map M [--apply]  như check/apply cho file html (chữ hiển thị + chuỗi trong <script>)
  fonts   <pptx> --min-pt 14 [--map M]

Map YAML: `entries` (pptx), `html` (html), `whitelist` (regex toàn cục), `fonts_allow`, `include` (glob tương đối).
Mục số: {slide: b07|7, in: body|notes, shape: id, text: "{} tin", keys: [file.json:path], fmt: int}.
Mục cho phép: {slide, allow: "lý do", shape?, text?}. Chart: {slide, chart: id, categories, series: [{name, key}]}.
"""
from __future__ import annotations

import argparse
import glob
import re
import sys
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.chart.data import CategoryChartData

sys.path.insert(0, str(Path(__file__).resolve().parent))
from slide_text import NUMBER, charts, frames, html_numbers  # noqa: E402
from slide_values import KeyBroken, Resolver, fmt  # noqa: E402

SLOT = r"([+\u2212-]?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?)"


def load_map(path: Path) -> dict:
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    for pattern in data.pop("include", []) or []:
        for sub in sorted(glob.glob(str(path.parent / pattern))):
            part = load_map(Path(sub))
            for k in ("entries", "html", "whitelist", "fonts_allow"):
                data.setdefault(k, []).extend(part.get(k, []) or [])
    return data


def compile_tpl(text: str) -> re.Pattern:
    out = []
    for piece in re.split(r"(\{\}|\{\*\})", text):
        out.append(SLOT if piece == "{}" else r"[\s\S]*?" if piece == "{*}" else re.escape(piece).replace(r"\ ", r"\s+"))
    return re.compile("".join(out))


def expected(entry: dict, res: Resolver) -> list[str]:
    keys = entry.get("keys") or [entry["key"]]
    fmts = entry.get("fmt", "int")
    fmts = fmts if isinstance(fmts, list) else [fmts] * len(keys)
    return [fmt(res.get(k), f) for k, f in zip(keys, fmts)]


def _slide_ok(entry, fr) -> bool:
    s = entry.get("slide")
    return s is None or s == fr.slide_name or s == fr.slide


def _scope(entry, fr) -> bool:
    where = entry.get("in", "any")
    return _slide_ok(entry, fr) and where in ("any", fr.where) and entry.get("shape") in (None, fr.shape_id)


class Report:
    def __init__(self, only: set[str] | None):
        self.only, self.fails = only, {"unmapped": 0, "broken": 0, "wrong": 0, "notfound": 0}

    def fail(self, kind: str, where: str, msg: str, slide_name: str | None = None) -> None:
        if self.only and slide_name not in self.only:
            return
        self.fails[kind] += 1
        print(f"FAIL {kind:8} {where}: {msg}")


def run_entries(items, texts, res, rep, apply_fn=None):
    """items: mục map; texts: [(label, slide_name, frame_like, text)] — trả tập (id(frame), start, end) đã phủ."""
    covered = set()
    for e in items:
        tpl = compile_tpl(e["text"]) if "text" in e else None
        label = f"slide {e.get('slide', '-')} '{e.get('text', e.get('shape', ''))}'"
        want = None
        if "allow" not in e:
            try:
                want = expected(e, res)
            except (KeyBroken, ValueError) as exc:
                rep.fail("broken", label, str(exc), e.get("slide"))
                continue
        hits = 0
        for lbl, name, fr, _ in texts:
            text = fr.text
            if not e.get("_match")(fr):
                continue
            if tpl is None:
                covered.add((id(fr), 0, len(text)))
                hits += 1
                continue
            found = list(tpl.finditer(text))
            for m in (reversed(found) if apply_fn else found):
                hits += 1
                if want is None:
                    covered.add((id(fr), m.start(), m.end()))
                    continue
                slots = [(m.start(i + 1), m.end(i + 1), m.group(i + 1)) for i in range(len(want))]
                for (a, b, got), w in zip(reversed(slots), reversed(want)):
                    if got != w:
                        if apply_fn:
                            apply_fn(fr, a, b, w)
                        else:
                            rep.fail("wrong", lbl, f"'{got}' ≠ '{w}' ({e.get('keys') or e.get('key')})", name)
                    covered.add((id(fr), a, a + len(w) if apply_fn else b))
        if not hits:
            rep.fail("notfound", label, "mẫu không khớp chữ nào", e.get("slide"))
    return covered


def _covered(spans, fid, a, b, wl_spans):
    return any(f == fid and s <= a and b <= t for f, s, t in spans) or any(s <= a and b <= t for s, t in wl_spans)


def check_pptx(prs, m, res, rep, apply=False):
    frs = frames(prs)
    entries = [e for e in m.get("entries", []) if "chart" not in e]
    for e in entries:
        e["_match"] = (lambda e: lambda fr: _scope(e, fr))(e)
    texts = [(f"slide {fr.slide}({fr.slide_name}) {fr.where} shape {fr.shape_id}", fr.slide_name, fr, fr.text) for fr in frs]
    spans = run_entries(entries, texts, res, rep, (lambda fr, a, b, w: fr.replace(a, b, w)) if apply else None)
    if apply:
        for idx, name, sid, chart in charts(prs):
            e = next((x for x in m.get("entries", []) if x.get("chart") == sid and x.get("slide") in (name, idx)), None)
            if e and "allow" not in e:
                check_chart(chart, e, res, rep, f"slide {idx} chart {sid}", name, True)
        return
    wl = [re.compile(w["re"]) for w in m.get("whitelist", [])]
    for lbl, name, fr, _ in texts:
        text = fr.text
        wl_spans = [(x.start(), x.end()) for r in wl for x in r.finditer(text)]
        for t in NUMBER.finditer(text):
            if not _covered(spans, id(fr), t.start(), t.end(), wl_spans):
                ctx = text[max(0, t.start() - 30):t.end() + 30].replace("\n", " ⏎ ").replace("\v", " ↵ ")
                rep.fail("unmapped", lbl, f"{t.group(0)!r} … {ctx!r}", name)
    chart_entries = {(e["slide"], e["chart"]): e for e in m.get("entries", []) if "chart" in e}
    for idx, name, sid, chart in charts(prs):
        e = chart_entries.get((name, sid)) or chart_entries.get((idx, sid))
        where = f"slide {idx}({name}) chart {sid}"
        if e is None:
            rep.fail("unmapped", where, "chart không có mục map", name)
        elif "allow" not in e:
            check_chart(chart, e, res, rep, where, name, apply)


def check_chart(chart, e, res, rep, where, name, apply):
    try:
        cats = e["categories"] if isinstance(e["categories"], list) else res.get(e["categories"])
        series = [(s["name"], [float(v) * s.get("scale", 1) for v in
                               ([res.get(k) for k in s["key"]] if isinstance(s["key"], list) else res.get(s["key"]))])
                  for s in e["series"]]
    except (KeyBroken, ValueError, TypeError) as exc:
        rep.fail("broken", where, str(exc), name)
        return
    if apply:
        data = CategoryChartData()
        data.categories = [str(c) for c in cats]
        for sname, vals in series:
            data.add_series(sname, vals)
        chart.replace_data(data)
        return
    got_cats = [str(c) for c in chart.plots[0].categories]
    got = [(se.name, list(se.values)) for pl in chart.plots for se in pl.series]
    if got_cats != [str(c) for c in cats]:
        rep.fail("wrong", where, f"nhãn {got_cats} ≠ {cats}", name)
    for (gn, gv), (wn, wv) in zip(got, series):
        if gn != wn or [round(v, 6) for v in gv] != [round(v, 6) for v in wv]:
            rep.fail("wrong", where, f"series {gn}={gv} ≠ {wn}={wv}", name)
    if len(got) != len(series):
        rep.fail("wrong", where, f"{len(got)} series ≠ {len(series)} trong map", name)


def check_html(raw: str, m, res, rep, apply=False) -> str:
    class Doc:
        text = raw

    doc = Doc()
    items = m.get("html", [])
    for e in items:
        e["_match"] = lambda fr: True

    def put(fr, a, b, w):
        fr.text = fr.text[:a] + w + fr.text[b:]

    spans = run_entries(items, [("html", None, doc, raw)], res, rep, put if apply else None)
    if apply:
        return doc.text
    wl = [(x.start(), x.end()) for w in m.get("whitelist", []) for x in re.finditer(w["re"], raw)]
    for a, b, tok in html_numbers(raw):
        if not _covered(spans, id(doc), a, b, wl):
            line = raw.count("\n", 0, a) + 1
            rep.fail("unmapped", f"html dòng {line}", f"{tok!r} … {raw[max(0, a - 40):b + 30]!r}")
    return raw


def fonts(prs, min_pt: float, m: dict) -> int:
    allow = {(a["slide"], a.get("shape")) for a in m.get("fonts_allow", [])}
    small = unresolved = 0
    for fr in frames(prs):
        if fr.where == "notes" or (fr.slide_name, fr.shape_id) in allow or (fr.slide_name, None) in allow:
            continue
        for _, _, run in fr.segments:
            if run is None or not run.text.strip():
                continue
            size = run.font.size or run._parent.font.size
            if size is None:
                unresolved += 1
            elif size.pt < min_pt:
                small += 1
                print(f"SMALL {size.pt:>4}pt slide {fr.slide}({fr.slide_name}) shape {fr.shape_id}: {run.text[:60]!r}")
    print(f"{small} run < {min_pt}pt · {unresolved} run cỡ chữ kế thừa (không kiểm)")
    return small


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", choices=["extract", "check", "apply", "html", "fonts"])
    ap.add_argument("file")
    ap.add_argument("--map", type=Path)
    ap.add_argument("--out", type=Path)
    ap.add_argument("--slides")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--min-pt", type=float, default=14)
    ap.add_argument("--results", type=Path, default=None)
    a = ap.parse_args(argv)
    m = load_map(a.map) if a.map else {}
    res = Resolver(a.results or Path(m.get("results_dir") or _default_results()))
    rep = Report(set(a.slides.split(",")) if a.slides else None)
    if a.cmd == "html":
        raw = Path(a.file).read_text(encoding="utf-8")
        out = check_html(raw, m, res, rep, apply=a.apply)
        if a.apply:
            Path(a.out or a.file).write_text(out, encoding="utf-8")
            rep = Report(None)
            check_html(out, m, res, rep)
    else:
        prs = Presentation(a.file)
        if a.cmd == "extract":
            for fr in frames(prs):
                for t in NUMBER.finditer(fr.text):
                    ctx = fr.text[max(0, t.start() - 30):t.end() + 30].replace("\n", " ⏎ ").replace("\v", " ↵ ")
                    print(f"{fr.slide}\t{fr.slide_name}\t{fr.where}\t{fr.shape_id}\t{t.group(0)}\t{ctx}")
            for idx, name, sid, ch in charts(prs):
                print(f"{idx}\t{name}\tchart\t{sid}\t-\t{[list(s.values) for p in ch.plots for s in p.series]}")
            return 0
        if a.cmd == "fonts":
            return 1 if fonts(prs, a.min_pt, m) else 0
        check_pptx(prs, m, res, rep, apply=a.cmd == "apply")
        if a.cmd == "apply":
            prs.save(a.out or a.file)
            rep = Report(set(a.slides.split(",")) if a.slides else None)
            check_pptx(Presentation(a.out or a.file), m, res, rep)
    f = rep.fails
    print(f"{f['unmapped']} unmapped, {f['broken']} broken keys, {f['wrong']} wrong, {f['notfound']} not found")
    return 1 if any(f.values()) else 0


def _default_results() -> str:
    import os
    return str(Path(os.environ.get("VRE_DATA_ROOT", "D:/Vietnam-Real-Estates")) / "results")


if __name__ == "__main__":
    sys.exit(main())
