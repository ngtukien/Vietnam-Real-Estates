"""Đọc/ghi chữ trong pptx (thân, bảng, notes) và vùng chữ hiển thị trong html, giữ định dạng run.

Mỗi khung chữ được trải phẳng thành một chuỗi: các đoạn nối bằng "\n", ngắt dòng `<a:br>` thành "\v".
`Frame.replace(start, end, new)` ghi đè đúng các run chứa khoảng [start, end).
"""
from __future__ import annotations

import html
import re
from dataclasses import dataclass, field

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.text.text import _Run

# Số kiểu VN: 3.500.744 · 0,656 · 46 — không dính chữ cái hai bên (2PN, e5, gpt-oss-120b bị bỏ qua).
NUMBER = re.compile(r"(?<![\w.,@#…])[+−-]?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?(?![\w])")


@dataclass
class Frame:
    slide: int            # 1-based, thứ tự cuối
    slide_name: str       # cSld@name (b06, n-eda-1, …)
    where: str            # body | notes
    shape_id: int
    text_frame: object = field(repr=False)
    cell: tuple | None = None  # (hàng, cột) khi là ô bảng
    segments: list = field(default_factory=list, repr=False)  # (start, end, run|None)
    text: str = ""

    def build(self) -> Frame:
        parts, pos, self.segments = [], 0, []
        for pi, p in enumerate(self.text_frame.paragraphs):
            if pi:
                parts.append("\n")
                pos += 1
            for child in p._p:
                tag = child.tag.rsplit("}", 1)[-1]
                if tag == "r":
                    run = _Run(child, p)
                    t = run.text
                elif tag == "br":
                    run, t = None, "\v"
                elif tag == "fld":
                    run, t = None, "".join(x.text or "" for x in child.iter() if x.tag.endswith("}t"))
                else:
                    continue
                self.segments.append((pos, pos + len(t), run))
                parts.append(t)
                pos += len(t)
        self.text = "".join(parts)
        return self

    def replace(self, start: int, end: int, new: str) -> None:
        hit = [(s, e, r) for s, e, r in self.segments if s < end and e > start]
        if not hit or any(r is None for _, _, r in hit):
            raise ValueError(f"không ghi được khoảng [{start},{end}) ở slide {self.slide} shape {self.shape_id}")
        first = True
        for s, e, r in hit:
            t = r.text
            a, b = max(start, s) - s, min(end, e) - s
            r.text = t[:a] + (new if first else "") + t[b:]
            first = False
        self.build()


def slide_name(slide) -> str:
    return slide._element.cSld.get("name") or ""


def _shapes(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(sh.shapes)
        else:
            yield sh


def frames(prs) -> list[Frame]:
    out = []
    for i, s in enumerate(prs.slides, 1):
        name = slide_name(s)
        for sh in _shapes(s.shapes):
            if sh.has_text_frame:
                out.append(Frame(i, name, "body", sh.shape_id, sh.text_frame).build())
            if getattr(sh, "has_table", False) and sh.has_table:
                for r, row in enumerate(sh.table.rows):
                    for c, cell in enumerate(row.cells):
                        out.append(Frame(i, name, "body", sh.shape_id, cell.text_frame, (r, c)).build())
        if s.has_notes_slide:
            out.append(Frame(i, name, "notes", 0, s.notes_slide.notes_text_frame).build())
    return out


def charts(prs):
    """(slide, slide_name, shape_id, chart) cho mọi chart nhúng."""
    for i, s in enumerate(prs.slides, 1):
        for sh in _shapes(s.shapes):
            if getattr(sh, "has_chart", False) and sh.has_chart:
                yield i, slide_name(s), sh.shape_id, sh.chart


_SCRIPT = re.compile(r"<script\b[^>]*>(.*?)</script>", re.S | re.I)
_STYLE = re.compile(r"<style\b[^>]*>.*?</style>", re.S | re.I)
_STRLIT = re.compile(r"'(?:\\.|[^'\\\n])*'|\"(?:\\.|[^\"\\\n])*\"|`(?:\\.|[^`\\])*`")
_TAG = re.compile(r"<[^>]*>")


def html_regions(raw: str) -> list[tuple[int, int]]:
    """Khoảng [start, end) của chữ người đọc thấy: text ngoài thẻ, và chuỗi ký tự trong <script>."""
    blocked = [(m.start(), m.end()) for m in _STYLE.finditer(raw)]
    regions = []
    for m in _SCRIPT.finditer(raw):
        blocked.append((m.start(), m.end()))
        regions += [(m.start(1) + s.start() + 1, m.start(1) + s.end() - 1) for s in _STRLIT.finditer(m.group(1))]
    blocked.sort()
    pos = 0
    for a, b in blocked + [(len(raw), len(raw))]:
        chunk = raw[pos:a]
        last = 0
        for t in _TAG.finditer(chunk):
            if t.start() > last:
                regions.append((pos + last, pos + t.start()))
            last = t.end()
        if last < len(chunk):
            regions.append((pos + last, a))
        pos = max(pos, b)
    return sorted(r for r in regions if r[1] > r[0])


def html_numbers(raw: str) -> list[tuple[int, int, str]]:
    out = []
    for a, b in html_regions(raw):
        seg = raw[a:b]
        # thực thể HTML (&#8211;, &nbsp;) không phải số trên trang
        masked = re.sub(r"&#?\w+;", lambda m: " " * len(m.group(0)), seg)
        out += [(a + m.start(), a + m.end(), m.group(0)) for m in NUMBER.finditer(masked)]
    return out


def unescape(s: str) -> str:
    return html.unescape(s)
