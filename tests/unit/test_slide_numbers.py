"""Công cụ kiểm số trên slide: scripts/slide_numbers.py + slide_values.py."""
import json
import sys
from pathlib import Path

import pytest
import yaml
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))
import slide_numbers  # noqa: E402
from slide_values import KeyBroken, Resolver, fmt  # noqa: E402


@pytest.fixture()
def results(tmp_path):
    d = tmp_path / "results"
    d.mkdir()
    (d / "r.json").write_text(json.dumps({
        "kept": 2537920, "acc": 0.6561, "rows": [{"id": "A", "v": 3.5}, {"id": "B", "v": 1.25}],
        "months": ["T1", "T2"], "counts": [10, 20]}), encoding="utf-8")
    return d


def _deck(tmp_path, body: str, notes: str = "", chart: bool = False) -> Path:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s._element.cSld.set("name", "x")
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = body
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    if chart:
        data = CategoryChartData()
        data.categories = ["T1", "T2"]
        data.add_series("n", [1, 2])
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(4), Inches(3), data)
    path = tmp_path / "d.pptx"
    prs.save(path)
    return path


def _map(tmp_path, entries, html=None) -> Path:
    path = tmp_path / "m.yaml"
    path.write_text(yaml.safe_dump({"entries": entries, "html": html or []}, allow_unicode=True), encoding="utf-8")
    return path


def _run(*args) -> int:
    return slide_numbers.main([str(a) for a in args])


def test_format_vietnamese_numbers():
    assert fmt(2537920, "int") == "2.537.920"
    assert fmt(0.6561, "pct1") == "65,6"
    assert fmt(0.656, "dec3") == "0,656"
    assert fmt(3500744, "mil1") == "3,5"
    assert fmt(71.45, "dec1") == "71,5"  # làm tròn nửa lên, không phải banker's rounding
    assert fmt(-0.05, "spct1") == "−5,0"
    assert fmt(0.05, "spct1") == "+5,0"
    assert fmt(1234567890123456, "last6") == "123456"


def test_resolver_paths_selectors_and_expressions(results):
    r = Resolver(results)
    assert r.get("r.json:rows[id=B].v") == 1.25
    assert r.get("r.json:rows[*].v") == [3.5, 1.25]
    assert r.get("r.json:rows[-1].id") == "B"
    assert r.get("=({r.json:kept}+0)/{r.json:counts[1]}") == 2537920 / 20
    for bad in ("r.json:nope", "x.json:kept", "r.json:rows[id=Z].v", "r.json:rows[5]"):
        with pytest.raises(KeyBroken):
            r.get(bad)


def test_unmapped_number_fails(tmp_path, results):
    deck = _deck(tmp_path, "Còn 2.537.920 tin, độ chính xác 65,6%")
    m = _map(tmp_path, [{"slide": "x", "text": "Còn {} tin", "key": "r.json:kept"}])
    assert _run("check", deck, "--map", m, "--results", results) == 1  # 65,6 chưa map


def test_mapped_numbers_pass(tmp_path, results):
    deck = _deck(tmp_path, "Còn 2.537.920 tin, độ chính xác 65,6%", notes="Lặp lại: 2.537.920 tin.")
    m = _map(tmp_path, [{"slide": "x", "in": "body", "text": "Còn {} tin, độ chính xác {}%",
                         "keys": ["r.json:kept", "r.json:acc"], "fmt": ["int", "pct1"]},
                        {"slide": "x", "in": "notes", "text": "{} tin", "key": "r.json:kept"}])
    assert _run("check", deck, "--map", m, "--results", results) == 0


def test_broken_key_and_wrong_value_fail(tmp_path, results):
    deck = _deck(tmp_path, "Còn 2.537.921 tin")
    assert _run("check", deck, "--map", _map(tmp_path, [{"slide": "x", "text": "Còn {} tin", "key": "r.json:kep"}]),
                "--results", results) == 1
    assert _run("check", deck, "--map", _map(tmp_path, [{"slide": "x", "text": "Còn {} tin", "key": "r.json:kept"}]),
                "--results", results) == 1


def test_apply_then_check_passes(tmp_path, results):
    deck = _deck(tmp_path, "Còn 19.925 tin, độ chính xác 0%", chart=True)
    m = _map(tmp_path, [
        {"slide": "x", "text": "Còn {} tin, độ chính xác {}%", "keys": ["r.json:kept", "r.json:acc"],
         "fmt": ["int", "pct1"]},
        {"slide": "x", "chart": 3, "categories": "r.json:months", "series": [{"name": "n", "key": "r.json:counts"}]}])
    assert _run("check", deck, "--map", m, "--results", results) == 1
    assert _run("apply", deck, "--map", m, "--results", results) == 0
    assert "2.537.920" in Presentation(deck).slides[0].shapes[0].text_frame.text


def test_allow_entry_with_reason(tmp_path, results):
    deck = _deck(tmp_path, "Ví dụ: L = 2000, S = 400")
    m = _map(tmp_path, [{"slide": "x", "text": "L = 2000, S = 400", "allow": "ví dụ công thức"}])
    assert _run("check", deck, "--map", m, "--results", results) == 0


def test_html_check_and_apply(tmp_path, results):
    page = tmp_path / "p.html"
    page.write_text("<p>Còn 1 tin</p><script>const T={a:[1,2]};const s='Bỏ 9 tin';</script>", encoding="utf-8")
    m = _map(tmp_path, [], html=[{"text": "Còn {} tin", "key": "r.json:kept"},
                                 {"text": "const T={a:[{},{}]}", "keys": ["r.json:counts[0]", "r.json:counts[1]"]}])
    assert _run("html", page, "--map", m, "--results", results) == 1  # 'Bỏ 9 tin' chưa map
    m = _map(tmp_path, [], html=[{"text": "Còn {} tin", "key": "r.json:kept"},
                                 {"text": "const T={a:[{},{}]}", "keys": ["r.json:counts[0]", "r.json:counts[1]"]},
                                 {"text": "Bỏ 9 tin", "allow": "ví dụ"}])
    assert _run("html", page, "--map", m, "--results", results, "--apply") == 0
    assert "Còn 2.537.920 tin" in page.read_text(encoding="utf-8") and "[10,20]" in page.read_text(encoding="utf-8")
